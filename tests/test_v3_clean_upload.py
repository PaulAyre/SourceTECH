"""v3: the browser strips, the server never stores a raw file, the vendor never sees a valuation.

Fixtures are SYNTHETIC (invented people) and were produced by the real browser
stripper: ScrapeTECH/stripper/tools/write_outputs.js.
"""
import io
import json
import os
import sys
import tempfile
import time
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
FIX = Path(__file__).resolve().parent / 'fixtures'
SECRET = 'pytest-ref-token-secret-0123456789abcdef0123'
PII = ['jason', 'nguyen', 'gmail.com', 'wattle street', '0412 345 678', '123 456 782', 'electrician', 'siobhan', 'papadopoulos']
FORBIDDEN = ['valuation_summary', 'estimated_value', 'total_valuation', 'master_document', 'multiple', 'portfolioapp', '987654321', '987,654,321', 'renewal_comm', 'product_type'] # internal field names never go out; the insurer's own header names may


@pytest.fixture()
def ctx(monkeypatch):
    d = tempfile.mkdtemp(prefix='st-v3-')
    monkeypatch.setenv('UPLOADS_DIR', d)
    monkeypatch.setenv('DATABASE', os.path.join(d, 't.db'))
    monkeypatch.setenv('TEMP_DIR', os.path.join(d, 'tmp'))
    monkeypatch.setenv('REF_TOKEN_SECRET', SECRET)
    monkeypatch.delenv('SOURCETECH_UI', raising=False)
    for m in [m for m in sys.modules if m in ('app', 'v3_api', 'email_service')]:
        del sys.modules[m]
    sys.path.insert(0, str(ROOT))
    import app as st
    calls = {'batches': [], 'emails': []}

    def fake_batch(file_paths, vendor_name, deal_id=None, generator_name=None):
        calls['batches'].append({'files': [Path(p).name for p in file_paths], 'vendor': vendor_name, 'deal_id': deal_id, 'generator': generator_name})
        return False, {'error': 'pytest: PavTECH not called'}

    def fake_email(**kw):
        calls['emails'].append(kw)
        return True, {'id': 'pytest'}

    monkeypatch.setattr(st.pavtech, 'process_batch', fake_batch)
    monkeypatch.setattr(st, 'send_dm_notification', fake_email)
    db = st.get_db()
    db.execute("INSERT INTO vendors (url_code, vendor_name, dm_email, dm_name, hubspot_deal_id, status) VALUES ('V3TEST01', 'Pytest Wealth', 'dm@example.com', 'Dee Em', '287657057728', 'complete')")
    vid = db.execute("SELECT id FROM vendors WHERE url_code='V3TEST01'").fetchone()[0]
    db.execute("INSERT INTO submissions (vendor_id, file_count, pavtech_status, master_document_path, valuation_summary, reference) VALUES (?, 2, 'complete', '/data/master_SECRET.xlsx', ?, 'ST-OLD')",
               (vid, json.dumps({'estimated_value': 987654321, 'total_valuation': 987654321})))
    db.commit(); db.close()
    return st, st.app.test_client(), Path(d), calls


def post_clean(client, name, insurer='aia', client_id='cid-000001', report=None):
    data = {'file': (io.BytesIO((FIX / name).read_bytes()), name), 'insurer': insurer, 'client_id': client_id,
            'safe_report': json.dumps(report or {'rulesVersion': '0.1.0', 'notes': [{'code': 'hidden_sheets_removed', 'count': 1}, {'code': 'Robert"); DROP', 'count': 1}, {'code': 'free text $1,234,567'}]})}
    return client.post('/V3TEST01/clean-upload', data=data, content_type='multipart/form-data')


def all_text(path):
    return ' '.join(str(v) for ws in load_workbook(path).worksheets for r in ws.iter_rows(values_only=True) for v in r if v is not None).lower()


def files_on_disk(d):
    return sorted(p.name for p in (d / 'V3TEST01').glob('*')) if (d / 'V3TEST01').exists() else []


def test_stripped_file_is_kept_and_its_codes_get_the_second_hash(ctx):
    st, c, d, _ = ctx
    before = load_workbook(FIX / 'aia_1.xlsx').worksheets[0]
    r = post_clean(c, 'aia_1.xlsx')
    assert r.status_code == 200, r.data
    j = r.get_json()
    assert j['success'] and not j['duplicate']
    assert j['file']['name'] == 'aia_1.xlsx' and j['file']['insurer'] == 'aia' and j['file']['policy_count'] > 0
    assert [n['code'] for n in j['file']['notes']] == ['hidden_sheets_removed'], 'only plain note codes reach the vendor'
    assert files_on_disk(d) == ['aia_1.xlsx'], 'exactly one file: the clean one'
    stored = d / 'V3TEST01' / 'aia_1.xlsx'
    rows = list(load_workbook(stored).worksheets[0].iter_rows(values_only=True))
    h = next(i for i, row in enumerate(rows) if 'ref_first' in row)
    col = rows[h].index('ref_first')
    raw_rows = list(before.iter_rows(values_only=True))
    import ref_tokens
    assert rows[h + 1][col] == ref_tokens.server_hash(SECRET, raw_rows[h + 1][col])
    assert rows[h + 1][col] != raw_rows[h + 1][col], 'the browser code is never what is stored'
    text = all_text(stored)
    for s in PII:
        assert s not in text


def test_raw_csv_is_refused_without_touching_the_disk(ctx):
    st, c, d, _ = ctx
    r = post_clean(c, 'raw_tal.csv', insurer='tal')
    assert r.status_code == 415 and r.get_json()['code'] == 'not_stripped'
    assert files_on_disk(d) == []


def test_unstripped_xlsx_is_cleaned_on_the_server_and_flagged_not_blocked(ctx, caplog):
    st, c, d, _ = ctx
    r = post_clean(c, 'raw_aia.xlsx')
    assert r.status_code == 200, 'nothing stops the vendor'
    codes = [n['code'] for n in r.get_json()['file']['notes']]
    assert 'server_removed_personal_details' in codes
    assert any('STRIP ALERT' in m for m in caplog.messages), 'and it fails LOUDLY for us'
    assert files_on_disk(d) == ['aia_1.xlsx']
    text = all_text(d / 'V3TEST01' / 'aia_1.xlsx')
    for s in PII:
        assert s not in text, s
    assert 'a1000001' in text


def test_same_file_twice_is_one_file(ctx):
    st, c, d, _ = ctx
    a = post_clean(c, 'tal_1.xlsx', insurer='tal', client_id='cid-aaaaaa').get_json()
    b = post_clean(c, 'tal_1.xlsx', insurer='tal', client_id='cid-bbbbbb').get_json()
    assert b['duplicate'] and b['file']['id'] == a['file']['id'] and files_on_disk(d) == ['tal_1.xlsx']


def test_without_the_secret_no_codes_are_stored(ctx, monkeypatch, caplog):
    st, c, d, _ = ctx
    monkeypatch.delenv('REF_TOKEN_SECRET')
    assert post_clean(c, 'tal_1.xlsx', insurer='tal').status_code == 200
    heads = [v for v in next(load_workbook(d / 'V3TEST01' / 'tal_1.xlsx').worksheets[0].iter_rows(values_only=True))]
    assert 'ref_first' not in heads and 'Policy Number' in heads
    assert any('REF_TOKEN_SECRET' in m for m in caplog.messages)


def test_the_old_raw_upload_endpoint_is_closed(ctx):
    st, c, d, _ = ctx
    r = c.post('/V3TEST01/upload', data={'file': (io.BytesIO((FIX / 'raw_tal.csv').read_bytes()), 'raw.csv')}, content_type='multipart/form-data')
    assert r.status_code == 410 and files_on_disk(d) == []


def test_vendor_never_sees_valuation_data_on_any_v3_route(ctx):
    st, c, d, _ = ctx
    post_clean(c, 'aia_1.xlsx')
    bodies = [c.get('/V3TEST01/app-config').get_data(as_text=True), c.get('/V3TEST01/status').get_data(as_text=True),
              c.post('/V3TEST01/submit-intent', json={'client_ids': ['cid-000001']}).get_data(as_text=True),
              c.post('/V3TEST01/submit-finalize').get_data(as_text=True)]
    for body in bodies:
        low = body.lower()
        for bad in FORBIDDEN:
            assert bad.lower() not in low, (bad, body[:300])
    cfg = c.get('/V3TEST01/app-config').get_json()
    assert cfg['vendor_name'] == 'Pytest Wealth' and len(cfg['insurers']) >= 9 and cfg['app_key'].startswith('iplus-ref-v1')
    assert cfg['latest_submission']['state'] in ('processed', 'needs_attention', 'processing', 'pending')
    assert set(cfg['files'][0]) == {'id', 'name', 'insurer', 'insurer_name', 'policy_count', 'uploaded_at', 'client_id', 'notes'}


def wait_for(cond, secs=5):
    t = time.time()
    while time.time() - t < secs:
        if cond():
            return True
        time.sleep(0.05)
    return False


def test_submit_before_uploads_finish_then_runs_when_the_last_file_lands(ctx):
    st, c, d, calls = ctx
    r = c.post('/V3TEST01/submit-intent', json={'client_ids': ['cid-000001', 'cid-000002']})
    assert r.status_code == 200 and r.get_json()['reference'].startswith('ST-')
    assert calls['batches'] == [], 'submit pressed, nothing has arrived yet: keep waiting'
    post_clean(c, 'aia_1.xlsx', client_id='cid-000001')
    assert calls['batches'] == []
    post_clean(c, 'tal_1.xlsx', insurer='tal', client_id='cid-000002')
    assert wait_for(lambda: len(calls['batches']) == 1)
    b = calls['batches'][0]
    assert sorted(b['files']) == ['aia_1.xlsx', 'tal_1.xlsx']
    assert b['deal_id'] == '287657057728' and b['generator'] == 'SourceTECH', 'HubSpot deal id goes to PavTECH'
    c.post('/V3TEST01/submit-finalize')
    time.sleep(0.2)
    assert len(calls['batches']) == 1, 'finalize is idempotent'


def test_tab_closed_after_submit_runs_with_what_arrived_and_tells_the_dm(ctx):
    st, c, d, calls = ctx
    c.post('/V3TEST01/submit-intent', json={'client_ids': ['cid-000001', 'cid-000002', 'cid-000003']})
    post_clean(c, 'aia_1.xlsx', client_id='cid-000001')
    db = st.get_db(); db.execute("UPDATE submission_intents SET deadline_at = '2000-01-01T00:00:00'"); db.commit(); db.close()
    st.v3['sweep_overdue']()
    assert wait_for(lambda: len(calls['batches']) == 1 and len(calls['emails']) == 1)
    assert calls['batches'][0]['files'] == ['aia_1.xlsx']
    notes = ' '.join(calls['emails'][0].get('extra_notes') or [])
    assert '2 of 3 files' in notes
    st.v3['sweep_overdue']()
    time.sleep(0.2)
    assert len(calls['batches']) == 1, 'a second sweep (or another gunicorn worker) cannot run it twice'
