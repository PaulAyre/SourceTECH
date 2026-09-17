"""Server-side defence in depth for files the browser has already stripped.

The browser does the real work. This pass should find NOTHING. If it finds
something, the browser stripper missed it (or was bypassed), so:
  - the unsafe column or cells are blanked here,
  - the upload carries on (nothing stops the vendor doing the job),
  - the caller raises a loud alert so the stripper gets fixed.

It also applies the second, secret hash to the relationship token columns so the
browser codes are never stored. Reads the SAME pii_rules.json and names.json as
the browser. No AI. Requires openpyxl.
"""
import json
import os
import re
import unicodedata

from openpyxl import load_workbook

import ref_tokens as rt

_HERE = os.path.dirname(os.path.abspath(__file__))
_names_cache = None


def load_names(path=None):
    global _names_cache
    if path is None and _names_cache is not None:
        return _names_cache
    with open(path or os.path.join(_HERE, "..", "data", "names.json"), encoding="utf-8") as fh:
        j = json.load(fh)
    n = {k: set(j[k]) for k in ("givenStrong", "givenWeak", "familyStrong", "familyWeak")}
    if path is None:
        _names_cache = n
    return n


# ---- header classification (port of src/headers.js) --------------------------------------
def tokenize_header(text, rules):
    s = rt.trim_all(text)
    s = re.sub(r"([a-z])([A-Z])", r"\1 \2", s)
    s = re.sub(r"([A-Za-z])(\d)", r"\1 \2", s)
    s = re.sub(r"(\d)([A-Za-z])", r"\1 \2", s)
    out = []
    for t in re.sub(r"[^a-z0-9]+", " ", s.lower()).split():
        glued = rules["gluedHeaderWords"].get(t)
        out.extend(glued.split(" ") if glued else [t])
    return out


def _find_seq(tokens, seq):
    n = len(seq)
    for i in range(len(tokens) - n + 1):
        if tokens[i:i + n] == seq:
            return i
    return -1


def _any_seq(tokens, seqs):
    for s in sorted(seqs, key=len, reverse=True):
        if _find_seq(tokens, s) >= 0:
            return s
    return None


def classify_header(text, rules):
    """Returns ('protected', valueKind) | ('pii', evidence) | ('unknown', None)."""
    tokens = tokenize_header(text, rules)
    if not tokens:
        return "unknown", None
    for o in rules["protectedOverrides"]:
        if _any_seq(tokens, o["seqs"]):
            return "protected", o["valueKind"]
    rest = list(tokens)
    for p in sorted(rules["protectedPhrases"], key=len, reverse=True):
        i = _find_seq(rest, p)
        while i >= 0:
            rest[i:i + len(p)] = [chr(0)]
            i = _find_seq(rest, p)
    person = next(g for g in rules["pii"] if g["evidence"] == "person")
    for g in rules["pii"]:
        hit = _any_seq(rest, g["seqs"])
        if not hit:
            continue
        if g["evidence"] == "name" and hit in (["name"], ["names"]):
            if not _any_seq(rest, person["seqs"]) and any(t in rules["nameExemptQualifiers"] for t in tokens):
                continue
        return "pii", g["evidence"]
    for p in rules["protected"]:
        if _any_seq(tokens, p["seqs"]):
            return "protected", p["valueKind"]
    return "unknown", None


# ---- value tests (port of src/valuescan.js) ----------------------------------------------
_EMAIL = re.compile(r"[A-Za-z0-9._%+'-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+")
_PHONE = re.compile(r"(?:\+?61[\s-]?|\(0[2-478]\)[\s-]?|\b0[2-478][\s-])\d(?:[\s-]?\d){7,8}\b|\b04\d{2}[\s-]\d{3}[\s-]\d{3}\b|\b1[38]00[\s-]\d{3}[\s-]\d{3}\b")
_TOTALS = re.compile(r"^(grand\s*)?(sub\s*)?totals?\b|^sum\b|^count\b|^average\b", re.I)   # "Total" often sits in the policy number column
_DIACRITICS = re.compile("[" + chr(0x0300) + "-" + chr(0x036F) + "]")
_APOS = re.compile("[" + "".join(chr(c) for c in (0x27, 0x60, 0x2018, 0x2019, 0x02BC)) + "]")


def _address_re(rules):
    types = "|".join(rules["valueScan"]["streetTypes"])
    street = re.compile(r"\b\d{1,5}[A-Za-z]?(?:\s*[/-]\s*\d{1,5})?\s+(?:[A-Za-z'.-]+\s+){1,4}(?:%s)\b\.?" % types, re.I)
    pobox = re.compile(r"\b(?:p\.?\s?o\.?\s?box|gpo\s?box|locked\s?bag|private\s?bag)\s*\d+", re.I)
    return lambda s: bool(street.search(s) or pobox.search(s))


def make_name_test(names, rules):
    gs, gw, fs, fw = names["givenStrong"], names["givenWeak"], names["familyStrong"], names["familyWeak"]
    titles, suffixes = set(rules["titles"]), set(rules["suffixes"])

    def hit(cell):
        s = rt.trim_all(cell)
        if not s or len(s) > 60 or re.search(r"\d", s) or re.search(r"[@#$%*=<>{}\[\]|\\]", s):
            return False
        w = _APOS.sub("", _DIACRITICS.sub("", unicodedata.normalize("NFKD", s)).lower())
        w = [x for x in re.split(r"[^a-z]+", w) if x]
        if not w or len(w) > 5:
            return False
        name_words = strong = given = 0
        for x in w:
            if x in titles or x in suffixes or x == "and" or len(x) == 1:
                continue
            is_strong = x in gs or x in fs
            if not (is_strong or x in gw or x in fw):
                return False
            name_words += 1
            strong += 1 if is_strong else 0
            given += 1 if (x in gs or x in gw) else 0
        if name_words == 0:
            return False
        if name_words == 1:
            return strong == 1
        return strong >= 1 or given >= 1
    return hit


def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return rt.trim_all(v)


def recheck_and_rehash(path, app_key, secret, rules=None, names=None, save_to=None):
    """Re-check a browser-stripped workbook IN PLACE and re-hash its token columns.

    Returns {'findings': [...], 'tokens_rehashed': n, 'token_columns_dropped': bool, 'rows': n}.
    findings never contain cell values, only where and what kind.
    """
    rules = rules or rt.load_rules()
    names = names or load_names()
    name_hit = make_name_test(names, rules)
    addr_hit = _address_re(rules)
    token_cols = set(rules["tokens"]["columns"] + rules["tokens"]["ownerColumns"])
    hash_cols = {c for c in token_cols if not c.endswith("_confidence")}
    lenient = {"id", "number", "date", "code"}
    findings, rehashed, rows_total = [], 0, 0
    drop_tokens = not secret or len(secret) < 32

    wb = load_workbook(path)          # not read_only: we may blank cells
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            findings.append({"sheet": ws.title, "kind": "hidden_sheet", "reason": "hidden_sheet_survived", "count": 1})
            wb.remove(ws)
            continue
        if name_hit(ws.title) or _EMAIL.search(ws.title):
            findings.append({"sheet": "(renamed)", "kind": "sheet_name", "reason": "name", "count": 1})
            ws.title = "Sheet%d" % (wb.worksheets.index(ws) + 1)
        grid = list(ws.iter_rows())
        if not grid:
            continue
        # header = the row holding ref_first, else the row with the most recognised headers
        h, best = -1, 0
        for r, row in enumerate(grid[: rules["valueScan"]["maxHeaderSearchRows"]]):
            texts = [_text(c.value) for c in row]
            if "ref_first" in texts:
                h = r
                break
            known = sum(1 for t in texts if t and classify_header(t, rules)[0] != "unknown")
            if known > best and known >= 2:
                h, best = r, known
        if h < 0:
            findings.append({"sheet": ws.title, "kind": "sheet", "reason": "no_table_found", "count": 1})
            wb.remove(ws)
            continue
        # rows that actually carry data (blank spacer rows do not count)
        token_idx = {i for i, c in enumerate(grid[h]) if _text(c.value) in token_cols}
        # a real row has a policy number, when the file has that column (totals and spacers do not count)
        id_idx = next((i for i, c in enumerate(grid[h]) if classify_header(_text(c.value), rules) == ("protected", "id")), None)
        if id_idx is not None and any(_text(row[id_idx].value) for row in grid[h + 1:] if id_idx < len(row)):
            rows_total += sum(1 for row in grid[h + 1:] if id_idx < len(row) and _text(row[id_idx].value) and not _TOTALS.match(_text(row[id_idx].value)))
        else:
            rows_total += sum(1 for row in grid[h + 1:] if any(_text(c.value) for i, c in enumerate(row) if i not in token_idx))

        # rows above the header
        for row in grid[:h]:
            for c in row:
                t = _text(c.value)
                if t and (_EMAIL.search(t) or _PHONE.search(t) or addr_hit(t) or name_hit(t)):
                    c.value = None
                    findings.append({"sheet": ws.title, "kind": "pre_header_cell", "reason": "personal_detail", "count": 1})

        for ci, hc in enumerate(grid[h]):
            header = _text(hc.value)
            cells = [row[ci] for row in grid[h + 1:] if ci < len(row)]
            if header in token_cols:
                if drop_tokens:
                    hc.value = None
                    for c in cells:
                        c.value = None
                    continue
                if header in hash_cols:
                    for c in cells:
                        v = _text(c.value)
                        if not v:
                            continue
                        try:
                            c.value = ";".join(rt.server_hash(secret, p) for p in v.split(";"))
                            rehashed += 1
                        except ValueError:
                            c.value = None     # not a token: something put real text in a token column
                            findings.append({"sheet": ws.title, "kind": "token_cell", "reason": "not_a_token", "count": 1})
                continue

            cls, kind = classify_header(header, rules)
            texts = [_text(c.value) for c in cells]
            if header and cls != "protected" and (_EMAIL.search(header) or _PHONE.search(header) or name_hit(header)):
                cls, kind = "pii", "header_is_pii"
            if cls == "pii":
                hc.value = None
                for c in cells:
                    c.value = None
                findings.append({"sheet": ws.title, "kind": "column", "reason": kind, "count": sum(1 for t in texts if t)})
                continue
            if cls == "protected" and kind in lenient:
                bad = [i for i, t in enumerate(texts) if t and _EMAIL.search(t)]
            else:
                bad = [i for i, t in enumerate(texts) if t and (_EMAIL.search(t) or _PHONE.search(t) or addr_hit(t))]
            name_hits = [] if (cls == "protected" and kind == "adviser") else [i for i, t in enumerate(texts) if t and name_hit(t)]
            if cls == "unknown" and (bad or name_hits):
                hc.value = None
                for c in cells:
                    c.value = None
                findings.append({"sheet": ws.title, "kind": "column", "reason": "value_scan", "count": len(bad) + len(name_hits)})
                continue
            if cls == "protected" and kind not in lenient and kind != "adviser":
                filled = sum(1 for t in texts if t)
                p = rules["valueScan"]["protectedTextColumn"]
                if len(name_hits) >= p["minNameHitsToFlag"] and filled and len(name_hits) / filled >= p["nameRateToFlag"]:
                    bad = sorted(set(bad + name_hits))
            for i in bad:
                cells[i].value = None
            if bad:
                findings.append({"sheet": ws.title, "kind": "cells", "reason": "personal_detail_in_needed_column", "count": len(bad)})

    if not wb.worksheets:
        raise ValueError("no_table_found")
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    wb.properties.title = ""
    wb.properties.company = "" if hasattr(wb.properties, "company") else None
    wb.save(save_to or path)
    if drop_tokens:
        findings.append({"sheet": "*", "kind": "config", "reason": "REF_TOKEN_SECRET_missing_token_columns_dropped", "count": 1})
    return {"findings": findings, "tokens_rehashed": rehashed, "token_columns_dropped": drop_tokens, "rows": rows_total}
